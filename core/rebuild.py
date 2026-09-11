# -*- coding: utf-8 -*-
"""
Rebuild all 8 sheets from Feishu data, matching daily-report-generator format exactly.
Sheet order: Total对比 / Regional对比 / Campaign对比 /
             数据分析-SEA / 数据分析-美洲 / 数据分析-东欧 / 数据分析 / metro&wow
"""
import io
import re
from collections import defaultdict, OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule

from .feishu import DEFAULT_REMARK

# ── layout constants ───────────────────────────────────────────────────────
SHEET_DEFS = [
    ('数据分析-SEA',  ['SEA'],                                False),
    ('数据分析-美洲', ['LATAM', 'NA'],                        False),
    ('数据分析-东欧', ['EEU'],                                False),
    ('数据分析',      ['WEU', 'MENA', 'GLOBAL', 'SA', 'AF'], True),
]
CATEGORIES = ('Newinstall', 'Reattribution', 'Active')
KEY_COUNTRIES = {'DE', 'EG', 'IQ', 'SA', 'PK'}
METRO_WOW_CODES = [
    'metro', 'wow',
    'competitor_rblx', 'competitor_ff', 'competitor_mlbb',
    'competitor_bs', 'competitor_stu',
]
REGION_ORDER = ['WEU', 'EEU', 'SA', 'MENA', 'LATAM', 'NA', 'SEA', 'AF', 'GLOBAL']

UAC_RE = re.compile(r'UAC(\d+\.?\d*)', re.IGNORECASE)

# ── fill colors ────────────────────────────────────────────────────────────
FILL_BLUE     = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
FILL_DARK_RED = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
FILL_GREEN    = PatternFill(start_color='548235', end_color='548235', fill_type='solid')
FILL_ORANGE   = PatternFill(start_color='ED7D31', end_color='ED7D31', fill_type='solid')
TODAY_FILL    = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
TOTAL_FILL    = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
SECTION_FILL  = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
HIGHLIGHT_FILL = PatternFill(start_color='F4B084', end_color='F4B084', fill_type='solid')
# analysis/metro uses slightly different blue
ANALY_CAT_FILL    = PatternFill(start_color='446EC4', end_color='446EC4', fill_type='solid')
ANALY_REGION_FILL = PatternFill(start_color='D9E3F3', end_color='D9E3F3', fill_type='solid')

# ── borders ────────────────────────────────────────────────────────────────
_thin = Side(style='thin', color='D9D9D9')
THIN_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
THICK_BOTTOM = Border(
    left=_thin, right=_thin, top=_thin,
    bottom=Side(style='medium', color='000000'),
)

# ── fonts ──────────────────────────────────────────────────────────────────
HDR_FONT  = Font(name='等线', bold=True, size=11, color='FFFFFF')
DATA_FONT = Font(name='等线', size=11)
SEC_FONT  = Font(name='等线', bold=True, size=11)

# analysis / metro fonts
CAT_FONT    = Font(name='等线', bold=True, size=11, color='FFFFFF')
REG_FONT    = Font(name='等线', size=10)
CTY_FONT    = Font(name='等线', size=10)
ANALY_FONT  = Font(name='等线', size=10)
BOLD10_FONT = Font(name='等线', bold=True, size=10)

COL_A_ALIGN = Alignment(horizontal='center', vertical='center')
COL_B_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)

# ── change-cell colors & format ────────────────────────────────────────────
CHG_UP   = '68A490'
CHG_DOWN = 'D65532'
CHG_ZERO = 'F1C232'
TOTAL_CHANGE_NUM_FMT = '"▲ "+0.00%;"▼ "0.00%;"▬ "0.00%'

# ── key-country highlights (Total对比) ────────────────────────────────────
HIGHLIGHT_COUNTRIES = {
    'TR', 'UZ', 'PK', 'ID', 'MY', 'TH', 'DE', 'US', 'EG', 'IQ', 'SA',
    'TR_IOS', 'UZ_IOS', 'PK_IOS', 'ID_IOS', 'MY_IOS', 'TH_IOS',
    'DE_IOS', 'US_IOS', 'EG_IOS', 'IQ_IOS', 'SA_IOS',
}

# ── row-height constants (analysis sheets) ────────────────────────────────
LINE_H   = 15.5
PAD      = 10.0
MIN_H    = 15.0
WRAP_CAP = 100

# ── keyword helpers (analysis sheets) ─────────────────────────────────────
BID_BUDGET_KW  = ('出价', '预算', 'TROAS')
ACTION_KW      = ('上调', '下调', '提升', '调整', '刺激', '降低', '提高', '增加', '减少', '控制', '减低')
NEW_RESTART_KW = ('开启首日', '重启首日')
PAUSED_KW      = ('已在后台暂停', '已暂停')

def is_bid_budget(t):  return bool(t) and any(k in t for k in BID_BUDGET_KW) and any(k in t for k in ACTION_KW)
def is_new_restart(t): return bool(t) and any(k in t for k in NEW_RESTART_KW)
def is_paused(t):      return bool(t) and any(k in t for k in PAUSED_KW)
def is_real_ai(t):     return bool(t) and t.strip() != DEFAULT_REMARK


# ── numeric helpers ────────────────────────────────────────────────────────
def safe_div(a, b): return a / b if b else 0.0
def safe_change(new_v, old_v): return safe_div((new_v or 0) - (old_v or 0), old_v)
def pct_change(y, t): return (t - y) / y if y else (1.0 if t else 0.0)

def fmt_int(v):
    try: return str(int(round(float(v))))
    except: return '0'

def fmt_num(v, d=2):
    try:
        v = float(v)
        if v == 0: return '0'
        if abs(v) >= 1 and v == int(v): return str(int(v))
        return f'{v:.{d}f}'
    except: return '0'

def cost_trend_word(c):   return '上升' if c > 0.10 else ('下降' if c < -0.10 else '稳定')
def volume_trend_word(c): return '增加' if c > 0.10 else ('下降' if c < -0.10 else '持平')
def cpi_trend_word(c):    return '上涨' if c > 0.10 else ('下降' if c < -0.10 else '稳定')
def uac_version(name):
    m = UAC_RE.search(name or '')
    return m.group(1) if m else '?'


# ── aggregation ────────────────────────────────────────────────────────────
def sum_metrics(camps):
    cost_y    = sum(c['cost_y']    for c in camps)
    cost_t    = sum(c['cost_t']    for c in camps)
    budget    = sum(c['budget']    for c in camps)
    clicks_y  = sum(c['clicks_y']  for c in camps)
    clicks_t  = sum(c['clicks_t']  for c in camps)
    install_y = sum(c['install_y'] for c in camps)
    install_t = sum(c['install_t'] for c in camps)
    uu_y      = sum(c.get('uu_y',     0) for c in camps)
    uu_t      = sum(c.get('uu_t',     0) for c in camps)
    impress_y = sum(c.get('impress_y',0) for c in camps)
    impress_t = sum(c.get('impress_t',0) for c in camps)
    reg_y     = sum(c.get('reg_y',    0) for c in camps)
    reg_t     = sum(c.get('reg_t',    0) for c in camps)
    login_y   = sum(c.get('login_y',  0) for c in camps)
    login_t   = sum(c.get('login_t',  0) for c in camps)
    return dict(
        cost_y=cost_y, cost_t=cost_t, budget=budget,
        clicks_y=clicks_y, clicks_t=clicks_t,
        install_y=install_y, install_t=install_t,
        uu_y=uu_y, uu_t=uu_t,
        impress_y=impress_y, impress_t=impress_t,
        reg_y=reg_y, reg_t=reg_t,
        login_y=login_y, login_t=login_t,
        cpi_y=safe_div(cost_y, install_y),
        cpi_t=safe_div(cost_t, install_t),
        cpc_y=safe_div(cost_y, clicks_y),
        cpc_t=safe_div(cost_t, clicks_t),
        cpr_reg_y=safe_div(cost_y, reg_y),
        cpr_reg_t=safe_div(cost_t, reg_t),
        cpr_login_y=safe_div(cost_y, login_y),
        cpr_login_t=safe_div(cost_t, login_t),
        cvr_y=safe_div(install_y, clicks_y),
        cvr_t=safe_div(install_t, clicks_t),
        freq_y=safe_div(impress_y, uu_y) if uu_y else 0.0,
        freq_t=safe_div(impress_t, uu_t) if uu_t else 0.0,
    )


# ── style helpers ──────────────────────────────────────────────────────────
_DARK_RED_HDRS = {'Cost', 'CPI', 'CPR(login)'}
_GREEN_HDRS    = {'Install', 'Registration'}
_ORANGE_HDRS   = {'Unique_users', 'impressions'}

def _hdr_fill(name, orange_extras=None):
    extras = set(orange_extras) if orange_extras else set()
    if name in _DARK_RED_HDRS:   return FILL_DARK_RED
    if name in _GREEN_HDRS:      return FILL_GREEN
    if name in (_ORANGE_HDRS | extras): return FILL_ORANGE
    return FILL_BLUE

def apply_header_row(ws, row, max_col, headers, orange_extras=None):
    current_group = None
    col_fills = {}
    for c, val in enumerate(headers):
        if val and val not in ('', '昨天', '今天', 'Change'):
            current_group = val
        col_fills[c] = _hdr_fill(current_group, orange_extras) if current_group else FILL_BLUE
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        cell.font      = HDR_FONT
        cell.fill      = col_fills.get(c - 1, FILL_BLUE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = THIN_BORDER
    ws.row_dimensions[row].height = 13.9

def apply_data_row(ws, row, max_col, today_cols=None):
    tc = today_cols or set()
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        cell.font      = DATA_FONT
        cell.border    = THIN_BORDER
        cell.alignment = Alignment(vertical='center')
        if c in tc:
            cell.fill = TODAY_FILL

def write_change_cell(ws, row, col, value):
    cell = ws.cell(row, col)
    if value is None or value == '-' or value == '':
        cell.value     = '#DIV/0!'
        cell.font      = Font(name='等线', size=11, color='000000')
        cell.alignment = Alignment(horizontal='distributed', vertical='center')
        cell.border    = THIN_BORDER
        return
    try:
        v = float(value)
    except (TypeError, ValueError):
        cell.value     = '#DIV/0!'
        cell.font      = Font(name='等线', size=11, color='000000')
        cell.alignment = Alignment(horizontal='distributed', vertical='center')
        cell.border    = THIN_BORDER
        return
    pct = abs(v) * 100
    if v > 0:
        icon, color, sign = '▲', CHG_UP, '+'
        pct_str = f' +{pct:.2f}%'
    elif v < 0:
        icon, color, sign = '▼', CHG_DOWN, '-'
        pct_str = f' -{pct:.2f}%'
    else:
        icon, color, sign = '▬', CHG_ZERO, ''
        pct_str = ' 0.00%'
    rt = CellRichText(
        TextBlock(InlineFont(rFont='等线', sz=11, color=color), icon),
        TextBlock(InlineFont(rFont='等线', sz=11, color='000000'), pct_str),
    )
    cell.value     = rt
    cell.alignment = Alignment(horizontal='distributed', vertical='center')
    cell.border    = THIN_BORDER

def _add_total_cond_fmt(ws, cell_ref):
    ws.conditional_formatting.add(cell_ref,
        CellIsRule(operator='greaterThan', formula=['0'],
                   font=Font(name='等线', size=11, color=CHG_UP)))
    ws.conditional_formatting.add(cell_ref,
        CellIsRule(operator='lessThan',    formula=['0'],
                   font=Font(name='等线', size=11, color=CHG_DOWN)))
    ws.conditional_formatting.add(cell_ref,
        CellIsRule(operator='equal',       formula=['0'],
                   font=Font(name='等线', size=11, color=CHG_ZERO)))

def write_subtotal_formula(ws, row, col, start, end):
    cl = get_column_letter(col)
    ws.cell(row, col, f'=SUBTOTAL(9,{cl}{start}:{cl}{end})')
    ws.cell(row, col).font          = DATA_FONT
    ws.cell(row, col).number_format = '#,##0.00'
    ws.cell(row, col).border        = THIN_BORDER

def write_change_formula(ws, row, col, t_col, y_col, start, end, derived=False):
    tc = get_column_letter(t_col)
    yc = get_column_letter(y_col)
    if derived:
        formula = f'=IFERROR(({tc}{row}-{yc}{row})/{yc}{row},0)'
    else:
        formula = (f'=IFERROR((SUBTOTAL(9,{tc}{start}:{tc}{end})'
                   f'-SUBTOTAL(9,{yc}{start}:{yc}{end}))'
                   f'/SUBTOTAL(9,{yc}{start}:{yc}{end}),0)')
    cell = ws.cell(row, col, formula)
    cell.number_format = TOTAL_CHANGE_NUM_FMT
    cell.alignment     = Alignment(horizontal='center')
    cell.border        = THIN_BORDER
    _add_total_cond_fmt(ws, f'{get_column_letter(col)}{row}')


# ── ① Campaign对比 ─────────────────────────────────────────────────────────
_CAMP_HDRS = [
    'Region', 'Country', 'Category', 'Campaign Code', 'Campaign',
    'Budget', '',                              # 6,7
    'Cost', 'Cost', 'Change',                  # 8,9,10
    'Clicks', 'Clicks', 'Change',              # 11,12,13
    'Install', 'Install', 'Change',            # 14,15,16
    'CPC', 'CPC', 'Change',                    # 17,18,19
    'CPI', 'CPI', 'Change',                    # 20,21,22
    'CPR(注册)', 'CPR(注册)', 'Change',         # 23,24,25
    'CPR(login)', 'CPR(login)', 'Change',       # 26,27,28
    'CVR', 'CVR', 'Change',                    # 29,30,31
    'Unique_users', 'Unique_users', 'Change',  # 32,33,34
    'Frequency', 'Frequency', 'Change',        # 35,36,37
    'impressions', 'impressions', 'Change',    # 38,39,40
    'Registration', 'Registration', 'Change',  # 41,42,43
    'Login_completed', 'Login_completed', 'Change',  # 44,45,46
]
_CAMP_MAX = len(_CAMP_HDRS)  # 46
_CAMP_TODAY  = {9, 12, 15, 18, 21, 24, 27, 30, 33, 36, 39, 42, 45}
_CAMP_CHANGE = [10, 13, 16, 19, 22, 25, 28, 31, 34, 37, 40, 43, 46]
# number format groups (1-indexed cols)
_CAMP_FMT_2DP = {8, 9, 17, 18, 20, 21, 23, 24, 26, 27}
_CAMP_FMT_INT = {11, 12, 14, 15, 32, 33, 38, 39, 41, 42, 44, 45}
_CAMP_FMT_PCT = {29, 30}  # CVR


def build_campaign_compare(ws, all_campaigns):
    max_col = _CAMP_MAX

    # header
    for c, v in enumerate(_CAMP_HDRS, 1):
        ws.cell(1, c, v)
    apply_header_row(ws, 1, max_col, _CAMP_HDRS)

    row = 2
    for camp in all_campaigns:
        m = sum_metrics([camp])
        row_data = [
            camp['a_region'], camp['b_country'], camp['c_category'],
            camp['d_code'], camp['e_name'],
            round(camp['budget']), round(camp.get('budget_t', camp['budget'])),
            round(m['cost_y'],  2), round(m['cost_t'],  2), safe_change(m['cost_t'],  m['cost_y']),
            round(m['clicks_y']),   round(m['clicks_t']),   safe_change(m['clicks_t'], m['clicks_y']),
            round(m['install_y']),  round(m['install_t']),  safe_change(m['install_t'],m['install_y']),
            round(m['cpc_y'],   4), round(m['cpc_t'],   4), safe_change(m['cpc_t'],   m['cpc_y']),
            round(m['cpi_y'],   4), round(m['cpi_t'],   4), safe_change(m['cpi_t'],   m['cpi_y']),
            round(m['cpr_reg_y'],4),round(m['cpr_reg_t'],4),safe_change(m['cpr_reg_t'],m['cpr_reg_y']),
            round(m['cpr_login_y'],4),round(m['cpr_login_t'],4),safe_change(m['cpr_login_t'],m['cpr_login_y']),
            round(m['cvr_y'],   6), round(m['cvr_t'],   6), safe_change(m['cvr_t'],   m['cvr_y']),
            round(m['uu_y']),       round(m['uu_t']),        safe_change(m['uu_t'],    m['uu_y']),
            round(m['freq_y'],  4), round(m['freq_t'],  4),  safe_change(m['freq_t'],  m['freq_y']),
            round(m['impress_y']),  round(m['impress_t']),   safe_change(m['impress_t'],m['impress_y']),
            round(m['reg_y']),      round(m['reg_t']),        safe_change(m['reg_t'],   m['reg_y']),
            round(m['login_y']),    round(m['login_t']),      safe_change(m['login_t'], m['login_y']),
        ]
        for c, v in enumerate(row_data, 1):
            ws.cell(row, c, v)

        apply_data_row(ws, row, max_col, _CAMP_TODAY)
        for col in _CAMP_CHANGE:
            if col <= max_col:
                write_change_cell(ws, row, col, ws.cell(row, col).value)
        for col in _CAMP_FMT_2DP:
            if col <= max_col: ws.cell(row, col).number_format = '#,##0.00'
        for col in _CAMP_FMT_INT:
            if col <= max_col: ws.cell(row, col).number_format = '#,##0'
        for col in _CAMP_FMT_PCT:
            if col <= max_col: ws.cell(row, col).number_format = '0.00%'
        row += 1

    # SUBTOTAL row
    data_start, data_end = 2, row - 1
    if data_end >= data_start:
        ws.cell(row, 5, 'Total').font = DATA_FONT
        for col in [6, 7, 8, 9, 11, 12, 14, 15, 32, 33, 38, 39, 41, 42, 44, 45]:
            if col <= max_col:
                write_subtotal_formula(ws, row, col, data_start, data_end)
        # derived: CPC, CPI, CPR, CVR, Freq
        for tgt, num, den in [
            (17,8,11),(18,9,12),(20,8,14),(21,9,15),
            (23,8,41),(24,9,42),(26,8,44),(27,9,45),
            (29,14,11),(30,15,12),(35,38,32),(36,39,33),
        ]:
            if tgt <= max_col:
                nc, dc = get_column_letter(num), get_column_letter(den)
                ws.cell(row, tgt, f'=IFERROR({nc}{row}/{dc}{row},0)')
                ws.cell(row, tgt).font = DATA_FONT
                ws.cell(row, tgt).number_format = '#,##0.00'
                ws.cell(row, tgt).border = THIN_BORDER
        for col in [29, 30]:
            if col <= max_col: ws.cell(row, col).number_format = '0.00%'
        # change formulas
        for chg, t_c, y_c in [(10,9,8),(13,12,11),(16,15,14),(34,33,32),(40,39,38),(43,42,41),(46,45,44)]:
            if chg <= max_col:
                write_change_formula(ws, row, chg, t_c, y_c, data_start, data_end, derived=False)
        for chg, t_c, y_c in [(19,18,17),(22,21,20),(25,24,23),(28,27,26),(31,30,29),(37,36,35)]:
            if chg <= max_col:
                write_change_formula(ws, row, chg, t_c, y_c, data_start, data_end, derived=True)
        for c in range(1, max_col + 1):
            ws.cell(row, c).fill   = TOTAL_FILL
            ws.cell(row, c).border = THIN_BORDER

    # widths / hide / freeze
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 55
    for c in range(6, max_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions['D'].hidden = True
    ws.column_dimensions['G'].hidden = True
    ws.sheet_format.defaultRowHeight = 15.0
    ws.sheet_format.customHeight = True
    ws.freeze_panes = 'F2'


# ── ② Total对比 ────────────────────────────────────────────────────────────
_TOTAL_HDRS = [
    'Region', 'Country', 'Budget', '消耗进度',           # 1,2,3,4
    'Cost', 'Cost', 'Change',                           # 5,6,7
    'Clicks', 'Clicks', 'Change',                       # 8,9,10
    'Install', 'Install', 'Change',                     # 11,12,13
    'CPC', 'CPC', 'Change',                             # 14,15,16
    'CPI', 'CPI', 'Change',                             # 17,18,19
    'CVR', 'CVR', 'Change',                             # 20,21,22
    'Unique_users', 'Unique_users', 'Change',            # 23,24,25
    'Frequency', 'Frequency', 'Change',                  # 26,27,28
    'impressions', 'impressions', 'Change',              # 29,30,31
    'Registration', 'Registration', 'Change',            # 32,33,34
    'CPR(注册)', 'CPR(注册)', 'Change',                  # 35,36,37
    'Login_completed', 'Login_completed', 'Change',      # 38,39,40
    'CPR(login)', 'CPR(login)', 'Change',                # 41,42,43
]
_TOTAL_MAX    = len(_TOTAL_HDRS)  # 43
_TOTAL_TODAY  = {6, 9, 12, 15, 18, 21, 24, 27, 30, 33, 36, 39, 42}
_TOTAL_CHANGE = [7, 10, 13, 16, 19, 22, 25, 28, 31, 34, 37, 40, 43]


def _region_sort_key(region):
    try: return REGION_ORDER.index(region)
    except ValueError: return len(REGION_ORDER)


def build_total_compare(ws, all_campaigns):
    max_col = _TOTAL_MAX
    row = 1

    for cat in CATEGORIES:
        # aggregate by (a_region, b_country) for this category
        country_camps: dict = OrderedDict()
        for c in all_campaigns:
            if c['c_category'] != cat: continue
            key = (c['a_region'], c['b_country'])
            country_camps.setdefault(key, []).append(c)
        if not country_camps: continue

        # sort by region order, then country
        sorted_pairs = sorted(country_camps.items(),
                               key=lambda x: (_region_sort_key(x[0][0]), x[0][1]))

        # section header (merged, D6E4F0)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
        cell = ws.cell(row, 1, cat)
        cell.font = SEC_FONT; cell.fill = SECTION_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = THIN_BORDER
        row += 1

        # column headers
        for c_i, v in enumerate(_TOTAL_HDRS, 1):
            ws.cell(row, c_i, v)
        apply_header_row(ws, row, max_col, _TOTAL_HDRS,
                         orange_extras=['Budget', '消耗进度'])
        row += 1

        data_start = row
        prev_region = None

        for (a_region, b_country), camps in sorted_pairs:
            m = sum_metrics(camps)
            spend = safe_div(m['cost_t'], m['budget'])

            if prev_region and a_region != prev_region:
                prev_r = row - 1
                for c_i in range(1, max_col + 1):
                    ws.cell(prev_r, c_i).border = THICK_BOTTOM

            prev_region = a_region

            row_data = [
                a_region, b_country, round(m['budget'], 2), spend,
                round(m['cost_y'],2), round(m['cost_t'],2), safe_change(m['cost_t'],  m['cost_y']),
                round(m['clicks_y']), round(m['clicks_t']),  safe_change(m['clicks_t'],m['clicks_y']),
                round(m['install_y']),round(m['install_t']), safe_change(m['install_t'],m['install_y']),
                round(m['cpc_y'],4),  round(m['cpc_t'],4),   safe_change(m['cpc_t'],   m['cpc_y']),
                round(m['cpi_y'],4),  round(m['cpi_t'],4),   safe_change(m['cpi_t'],   m['cpi_y']),
                round(m['cvr_y'],6),  round(m['cvr_t'],6),   safe_change(m['cvr_t'],   m['cvr_y']),
                round(m['uu_y']),     round(m['uu_t']),       safe_change(m['uu_t'],    m['uu_y']),
                round(m['freq_y'],4), round(m['freq_t'],4),  safe_change(m['freq_t'],  m['freq_y']),
                round(m['impress_y']),round(m['impress_t']), safe_change(m['impress_t'],m['impress_y']),
                round(m['reg_y']),    round(m['reg_t']),      safe_change(m['reg_t'],   m['reg_y']),
                round(m['cpr_reg_y'],4),round(m['cpr_reg_t'],4), safe_change(m['cpr_reg_t'],m['cpr_reg_y']),
                round(m['login_y']),  round(m['login_t']),    safe_change(m['login_t'], m['login_y']),
                round(m['cpr_login_y'],4),round(m['cpr_login_t'],4),safe_change(m['cpr_login_t'],m['cpr_login_y']),
            ]
            for c_i, v in enumerate(row_data, 1):
                ws.cell(row, c_i, v)

            apply_data_row(ws, row, max_col, _TOTAL_TODAY)
            for col in _TOTAL_CHANGE:
                write_change_cell(ws, row, col, ws.cell(row, col).value)

            # number formats
            for col in [3, 5, 6]: ws.cell(row, col).number_format = '#,##0.00'
            ws.cell(row, 4).number_format = '0.0%'
            for col in [8, 9, 11, 12, 23, 24, 29, 30, 32, 33, 38, 39]:
                ws.cell(row, col).number_format = '#,##0'
            for col in [14, 15, 17, 18, 20, 21, 26, 27, 35, 36, 41, 42]:
                ws.cell(row, col).number_format = '#,##0.00'

            if b_country in HIGHLIGHT_COUNTRIES:
                ws.cell(row, 2).fill = HIGHLIGHT_FILL

            row += 1

        # SUBTOTAL row
        data_end = row - 1
        if data_end >= data_start:
            ws.cell(row, 1, 'Total').font = DATA_FONT
            for col in [3, 5, 6, 8, 9, 11, 12, 23, 24, 29, 30, 32, 33, 38, 39]:
                write_subtotal_formula(ws, row, col, data_start, data_end)
            # derived
            for tgt, num, den in [
                (14,5,8),(15,6,9),(17,5,11),(18,6,12),
                (20,11,8),(21,12,9),(26,29,23),(27,30,24),
                (35,5,32),(36,6,33),(41,5,38),(42,6,39),
            ]:
                nc, dc = get_column_letter(num), get_column_letter(den)
                ws.cell(row, tgt, f'=IFERROR({nc}{row}/{dc}{row},0)')
                ws.cell(row, tgt).font = DATA_FONT
                ws.cell(row, tgt).number_format = '#,##0.00'
                ws.cell(row, tgt).border = THIN_BORDER
            for col in [20, 21]: ws.cell(row, col).number_format = '0.00%'
            # spend rate
            cl, bl = get_column_letter(6), get_column_letter(3)
            ws.cell(row, 4, f'=IFERROR({cl}{row}/{bl}{row},0)')
            ws.cell(row, 4).font = DATA_FONT; ws.cell(row, 4).number_format = '0.0%'
            ws.cell(row, 4).border = THIN_BORDER
            # change formulas
            for chg, t_c, y_c in [(7,6,5),(10,9,8),(13,12,11),(25,24,23),(31,30,29),(34,33,32),(40,39,38)]:
                write_change_formula(ws, row, chg, t_c, y_c, data_start, data_end, derived=False)
            for chg, t_c, y_c in [(16,15,14),(19,18,17),(22,21,20),(28,27,26),(37,36,35),(43,42,41)]:
                write_change_formula(ws, row, chg, t_c, y_c, data_start, data_end, derived=True)
            for c_i in range(1, max_col + 1):
                ws.cell(row, c_i).fill = TOTAL_FILL
                ws.cell(row, c_i).border = THIN_BORDER
            row += 1

        row += 1  # blank row between sections

    # column widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 16
    for c in range(3, _TOTAL_MAX + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.sheet_format.defaultRowHeight = 15.0
    ws.sheet_format.customHeight = True

    # conditional: spend rate < 70% → red fill
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    ws.conditional_formatting.add(
        f'D2:D{row}',
        FormulaRule(formula=['AND(D2<>"",D2<0.7)'], fill=red_fill),
    )
    ws.freeze_panes = 'C3'


# ── ③ Regional对比 ─────────────────────────────────────────────────────────
_REG_HDRS = [
    'Region', 'Location', 'Country',
    'Cost', 'Cost', 'Change',      # 4,5,6
    'Clicks', 'Clicks', 'Change',  # 7,8,9
    'Install', 'Install', 'Change',# 10,11,12
    'CPC', 'CPC', 'Change',        # 13,14,15
    'CPI', 'CPI', 'Change',        # 16,17,18
    'CVR', 'CVR', 'Change',        # 19,20,21
]
_REG_MAX = len(_REG_HDRS)  # 21


def build_regional_compare(ws, all_campaigns):
    groups: dict = OrderedDict()
    for c in all_campaigns:
        uac = uac_version(c['e_name'])
        key = (c['a_region'], c['b_country'], uac)
        groups.setdefault(key, []).append(c)

    # header
    for c_i, v in enumerate(_REG_HDRS, 1):
        ws.cell(1, c_i, v)
    apply_header_row(ws, 1, _REG_MAX, _REG_HDRS)

    last_row = 1
    for r_i, ((a_region, b_country, uac), camps) in enumerate(groups.items(), 2):
        m = sum_metrics(camps)
        row_data = [
            a_region, b_country, f'{b_country}-{uac}',
            round(m['cost_y'],   6), round(m['cost_t'],   6), safe_change(m['cost_t'],  m['cost_y']),
            round(m['clicks_y']),    round(m['clicks_t']),     safe_change(m['clicks_t'],m['clicks_y']),
            round(m['install_y']),   round(m['install_t']),    safe_change(m['install_t'],m['install_y']),
            round(m['cpc_y'],    6), round(m['cpc_t'],    6),  safe_change(m['cpc_t'],   m['cpc_y']),
            round(m['cpi_y'],    6), round(m['cpi_t'],    6),  safe_change(m['cpi_t'],   m['cpi_y']),
            round(m['cvr_y'],    6), round(m['cvr_t'],    6),  safe_change(m['cvr_t'],   m['cvr_y']),
        ]
        for c_i, v in enumerate(row_data, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.font   = DATA_FONT
            cell.border = THIN_BORDER
        for col in [4, 5, 13, 14, 16, 17]: ws.cell(r_i, col).number_format = '#,##0.000000'
        for col in [7, 8, 10, 11]:          ws.cell(r_i, col).number_format = '#,##0'
        for col in [6, 9, 12, 15, 18, 21]:  ws.cell(r_i, col).number_format = '0.00%'
        for col in [19, 20]:                ws.cell(r_i, col).number_format = '0.00%'
        last_row = r_i

    # conditional formatting on change cols
    for col_i in [6, 9, 12, 15, 18, 21]:
        cl = get_column_letter(col_i)
        rng = f'{cl}2:{cl}{last_row}'
        ws.conditional_formatting.add(rng,
            CellIsRule(operator='greaterThan', formula=['0.1'],
                       fill=PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'),
                       font=Font(name='等线', size=11, color='006100')))
        ws.conditional_formatting.add(rng,
            CellIsRule(operator='lessThan', formula=['-0.1'],
                       fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
                       font=Font(name='等线', size=11, color='9C0006')))

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 22
    for c in range(4, _REG_MAX + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.sheet_format.defaultRowHeight = 15.0
    ws.sheet_format.customHeight = True
    ws.freeze_panes = 'D2'


# ── ④ metro&wow ────────────────────────────────────────────────────────────
def _mw_code(d_code: str):
    cl = d_code.lower()
    for k in METRO_WOW_CODES:
        if cl == k.lower():
            return k
    return None

def _make_overview_text(m):
    ct = cost_trend_word(pct_change(m['cost_y'], m['cost_t']))
    vt = volume_trend_word(pct_change(m['install_y'], m['install_t']))
    pt = cpi_trend_word(pct_change(m['cpi_y'], m['cpi_t']))
    return (f"整体花费{ct}{fmt_int(m['cost_y'])}-{fmt_int(m['cost_t'])}/{fmt_int(m['budget'])}，"
            f"量级{vt}{fmt_int(m['install_y'])}-{fmt_int(m['install_t'])}，"
            f"成本{pt}{fmt_num(m['cpi_y'])}-{fmt_num(m['cpi_t'])}，{DEFAULT_REMARK}")

def _make_country_text(code, m):
    ct = cost_trend_word(pct_change(m['cost_y'], m['cost_t']))
    vt = volume_trend_word(pct_change(m['install_y'], m['install_t']))
    pt = cpi_trend_word(pct_change(m['cpi_y'], m['cpi_t']))
    return (f"—{code}花费{ct}{fmt_int(m['cost_y'])}-{fmt_int(m['cost_t'])}/{fmt_int(m['budget'])}，"
            f"量级{vt}{fmt_int(m['install_y'])}-{fmt_int(m['install_t'])}，"
            f"成本{pt}{fmt_num(m['cpi_y'])}-{fmt_num(m['cpi_t'])}，{DEFAULT_REMARK}")


def build_metro_wow(ws, all_campaigns):
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 115

    # aggregate: code → country → [camps]
    sections: dict = {k: defaultdict(list) for k in METRO_WOW_CODES}
    for c in all_campaigns:
        key = _mw_code(c['d_code'])
        if key is None: continue
        sections[key][c['b_country']].append(c)

    out_row = [1]

    def wr(label_a, text_b, font_a, font_b):
        ca = ws.cell(out_row[0], 1, label_a)
        cb = ws.cell(out_row[0], 2, text_b)
        ca.font = font_a; ca.alignment = COL_A_ALIGN
        cb.font = font_b; cb.alignment = COL_B_ALIGN
        # row height based on line count
        lines = (text_b or '').count('\n') + 1
        ws.row_dimensions[out_row[0]].height = max(MIN_H, lines * LINE_H)
        out_row[0] += 1

    for code in METRO_WOW_CODES:
        country_map = sections[code]
        # filter active countries
        active = {bc: camps for bc, camps in country_map.items()
                  if any(c['cost_y'] > 0 or c['cost_t'] > 0 for c in camps)}
        if not active: continue

        # section header
        section_title = f'{code.upper()}-newinstall'
        row_idx = out_row[0]
        ca = ws.cell(row_idx, 1, section_title)
        cb = ws.cell(row_idx, 2, '')
        ca.font = HDR_FONT;  ca.fill = ANALY_CAT_FILL; ca.alignment = COL_A_ALIGN
        cb.font = HDR_FONT;  cb.fill = ANALY_CAT_FILL; cb.alignment = COL_B_ALIGN
        ws.row_dimensions[row_idx].height = MIN_H
        out_row[0] += 1

        # TOTAL row
        all_camps = [c for camps in active.values() for c in camps]
        total_m = sum_metrics(all_camps)
        wr('TOTAL', _make_overview_text(total_m), CTY_FONT, BOLD10_FONT)

        # country rows sorted by cost_t descending
        sorted_countries = sorted(
            active.items(),
            key=lambda x: sum(c['cost_t'] for c in x[1]),
            reverse=True,
        )
        for b_country, camps in sorted_countries:
            cty_m = sum_metrics(camps)
            wr(b_country, _make_country_text(code, cty_m), CTY_FONT, ANALY_FONT)


# ── ⑤ 数据分析 series ─────────────────────────────────────────────────────
def _make_ov_prefix(m, ios=False):
    leader = 'IOS整体' if ios else '整体'
    ct = cost_trend_word(pct_change(m['cost_y'], m['cost_t']))
    vt = volume_trend_word(pct_change(m['install_y'], m['install_t']))
    pt = cpi_trend_word(pct_change(m['cpi_y'], m['cpi_t']))
    return (f"{leader}花费{ct}{fmt_int(m['cost_y'])}-{fmt_int(m['cost_t'])}/{fmt_int(m['budget'])}，"
            f"量级{vt}{fmt_int(m['install_y'])}-{fmt_int(m['install_t'])}，"
            f"成本{pt}{fmt_num(m['cpi_y'])}-{fmt_num(m['cpi_t'])}，")

def _make_cp_prefix(camp):
    ct = cost_trend_word(pct_change(camp['cost_y'], camp['cost_t']))
    vt = volume_trend_word(pct_change(camp['install_y'], camp['install_t']))
    pt = cpi_trend_word(pct_change(camp['cpi_y'], camp['cpi_t']))
    return (f"-{camp['d_code']}花费{ct}{fmt_int(camp['cost_y'])}-{fmt_int(camp['cost_t'])}/{fmt_int(camp['budget'])}，"
            f"量级{vt}{fmt_int(camp['install_y'])}-{fmt_int(camp['install_t'])}，"
            f"成本{pt}{fmt_num(camp['cpi_y'])}-{fmt_num(camp['cpi_t'])}，")


def build_rich_cell(line_parts):
    if not any(bb or nr for _, _, bb, nr in line_parts):
        return '\n'.join((p or '') + (r or '') for p, r, b, red in line_parts)
    _reg = InlineFont(rFont='等线', sz=10)
    _red = InlineFont(rFont='等线', sz=10, color='FF0000')
    blocks = []
    for i, (p, r, is_bb, is_nr) in enumerate(line_parts):
        if i > 0:
            blocks.append('\n')
        if p:
            blocks.append(TextBlock(_reg, p))
        if r:
            blocks.append(TextBlock(_red if (is_bb or is_nr) else _reg, r))
    return CellRichText(blocks)

def _cell_text(v):
    if isinstance(v, CellRichText):
        return ''.join(blk.text if hasattr(blk, 'text') else str(blk) for blk in v)
    return str(v) if v is not None else ''


def _build_analysis_sheets(wb, all_campaigns, country_total_ai):
    grouped: dict = defaultdict(lambda: defaultdict(OrderedDict))
    for camp in all_campaigns:
        ar, cat, bc = camp['a_region'], camp['c_category'], camp['b_country']
        grouped[ar][cat].setdefault(bc, []).append(camp)

    for sheet_name, region_order, filter_structure in SHEET_DEFS:
        ws = wb.create_sheet(sheet_name)
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 115
        out_row = [1]

        def wr_header(label, font, fill):
            ca = ws.cell(out_row[0], 1, label); cb = ws.cell(out_row[0], 2, '')
            ca.font = font;  ca.fill = fill;  ca.alignment = COL_A_ALIGN
            cb.font = font;  cb.fill = fill;  cb.alignment = COL_B_ALIGN
            out_row[0] += 1

        def wr_row(label, parts):
            ca = ws.cell(out_row[0], 1, label); cb = ws.cell(out_row[0], 2)
            cb.value = build_rich_cell(parts)
            ca.font = CTY_FONT;  ca.alignment = COL_A_ALIGN
            cb.font = ANALY_FONT; cb.alignment = COL_B_ALIGN
            out_row[0] += 1

        for category in CATEGORIES:
            has_data = any(bool(grouped.get(ar, {}).get(category)) for ar in region_order)
            if not has_data: continue
            wr_header(category, CAT_FONT, ANALY_CAT_FILL)

            for a_region in region_order:
                if a_region not in grouped or category not in grouped[a_region]: continue
                country_map = grouped[a_region][category]
                android_bcs = [bc for bc in country_map if not bc.upper().endswith('_IOS')]
                all_android  = [c for bc in android_bcs for c in country_map[bc]]
                if not all_android: continue

                wr_header(a_region, REG_FONT, ANALY_REGION_FILL)

                # region TOTAL
                rtm = sum_metrics(all_android)
                if rtm['cost_y'] != 0 or rtm['cost_t'] != 0:
                    wr_row('TOTAL', [(_make_ov_prefix(rtm), DEFAULT_REMARK, False, False)])

                for bc in android_bcs:
                    android_c = country_map[bc]
                    ios_b     = bc + '_IOS'
                    ios_c     = country_map.get(ios_b, [])

                    am = sum_metrics(android_c)
                    country_ai_raw = country_total_ai.get((a_region, bc, category), '')
                    country_ai = country_ai_raw if is_real_ai(country_ai_raw) else None

                    ai_for_ov = country_ai_raw or DEFAULT_REMARK
                    ov_part   = (_make_ov_prefix(am), ai_for_ov,
                                 is_bid_budget(ai_for_ov), is_new_restart(ai_for_ov))

                    ios_ov = None
                    if ios_c:
                        im = sum_metrics(ios_c)
                        ios_ai = country_total_ai.get((a_region, ios_b, category), '') or DEFAULT_REMARK
                        ios_ov = (_make_ov_prefix(im, ios=True), ios_ai,
                                  is_bid_budget(ios_ai), is_new_restart(ios_ai))

                    all_cp, abnormal_cp = [], []
                    for camp in android_c:
                        ai = camp['ai'] or DEFAULT_REMARK
                        if is_paused(ai): continue
                        part = (_make_cp_prefix(camp), ai, is_bid_budget(ai), is_new_restart(ai))
                        all_cp.append(part)
                        if is_real_ai(camp['ai']): abnormal_cp.append(part)

                    ios_cp = []
                    for camp in ios_c:
                        ai = camp['ai'] or DEFAULT_REMARK
                        if is_paused(ai): continue
                        if is_real_ai(camp['ai']):
                            ios_cp.append((_make_cp_prefix(camp), ai,
                                           is_bid_budget(ai), is_new_restart(ai)))

                    if filter_structure:
                        if not (bc in KEY_COUNTRIES or abnormal_cp or ios_c or country_ai):
                            continue
                        parts = [ov_part]
                        if ios_ov: parts.append(ios_ov)
                        # KEY_COUNTRIES: show ALL non-paused campaigns
                        if bc in KEY_COUNTRIES:
                            parts.extend(all_cp)
                        else:
                            parts.extend(abnormal_cp)
                        parts.extend(ios_cp)
                    else:
                        parts = [ov_part]
                        if ios_ov: parts.append(ios_ov)
                        parts.extend(all_cp)
                        parts.extend(ios_cp)

                    wr_row(bc, parts)


def fix_analysis_row_heights(wb, analysis_sheet_names):
    for ws in wb.worksheets:
        if ws.title not in analysis_sheet_names: continue
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=2).value
            if v is None: continue
            text = _cell_text(v)
            n = sum(
                max(1, -(-( sum(2 if ord(c) > 0x2E80 else 1 for c in ln) ) // WRAP_CAP))
                for ln in text.split('\n')
            )
            ws.row_dimensions[r].height = max(round(PAD + n * LINE_H, 2), MIN_H)


# ── main entry point ───────────────────────────────────────────────────────
def rebuild_from_feishu(all_campaigns: list, country_total_ai: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('Total对比')
    build_total_compare(ws, all_campaigns)

    ws = wb.create_sheet('Regional对比')
    build_regional_compare(ws, all_campaigns)

    ws = wb.create_sheet('Campaign对比')
    build_campaign_compare(ws, all_campaigns)

    _build_analysis_sheets(wb, all_campaigns, country_total_ai)

    ws = wb.create_sheet('metro&wow')
    build_metro_wow(ws, all_campaigns)

    analysis_names = {sd[0] for sd in SHEET_DEFS} | {'metro&wow'}
    fix_analysis_row_heights(wb, analysis_names)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.read()
