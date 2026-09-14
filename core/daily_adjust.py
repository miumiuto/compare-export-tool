# -*- coding: utf-8 -*-
import io
import openpyxl
from openpyxl.utils import get_column_letter

MAX_COL = 33  # A through AG

# Within-row derived formulas: col_idx → rhs template (use {r} for row number)
_DERIVED = {
    3:  'E{r}+M{r}+R{r}+V{r}+AB{r}',   # C: Total Spent
    4:  'H{r}+N{r}+S{r}+W{r}+AC{r}',   # D: Total Imp
    6:  'E{r}/G{r}',                     # F: CPI
    10: 'I{r}/H{r}',                     # J: CTR
    11: 'G{r}/I{r}',                     # K: CVR
    12: 'G{r}/H{r}',                     # L: IR
    16: 'M{r}/O{r}',                     # P: CPV
    17: 'M{r}/N{r}*1000',               # Q: CPM
    21: 'R{r}/T{r}',                     # U: CPC
    25: 'V{r}/X{r}',                     # Y: cpc (Reattr)
    26: 'V{r}/AA{r}',                    # Z: CPR (Reattr)
    31: 'AB{r}/AD{r}',                   # AE: cpc (Active)
    32: 'AB{r}/AG{r}',                   # AF: CPR (Active)
}


def _eeu_formula(col_idx, row):
    cl = get_column_letter(col_idx)
    r  = row
    if col_idx == 1: return f'=AZ!A{r}'
    if col_idx == 2: return f'=AZ!B{r}'
    if col_idx in _DERIVED: return '=' + _DERIVED[col_idx].format(r=r)

    az  = f"AZ!{cl}{r}"
    eeu = f"'EEU（KZ、KG、BY）'!{cl}{r}"
    eo  = f"EEU_Others!{cl}{r}"
    kz  = f"KZ!{cl}{r}"
    return f'={az}+{eeu}+{eo}+{kz}'


def _sa_formula(col_idx, row):
    cl = get_column_letter(col_idx)
    r  = row
    if col_idx == 1: return f'=SA!A{r}'
    if col_idx == 2: return f'=SA!B{r}'
    if col_idx in _DERIVED: return '=' + _DERIVED[col_idx].format(r=r)
    return f'=SA!{cl}{r}+SA_IOS!{cl}{r}'


def _fill_formulas(ws_new, formula_fn, ref_ws):
    """Replace data rows (3+) with formulas; stop at first empty Date cell."""
    for r in range(3, ref_ws.max_row + 1):
        if ref_ws.cell(r, 2).value is None:
            break
        for col_idx in range(1, MAX_COL + 1):
            ws_new.cell(r, col_idx).value = formula_fn(col_idx, r)


def process_eeu(file_bytes: bytes) -> bytes:
    """
    Idempotent EEU adjustment:
      - Copy 'EEU TOTAL' (or 'AZ' as fallback) to get formatting template
      - Delete old EEU其他 TOTAL if present
      - Rename new copy → 'EEU其他 TOTAL', update A1 → 'CIS TOTAL'
      - Rename 'EEU' → 'EEU（KZ、KG、BY）' if needed
      - Delete 'EEU TOTAL' if present
      - Fill data rows with cross-sheet formulas
      - Move 'EEU其他 TOTAL' to position 0
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))

    # Remove stale EEU其他 TOTAL (idempotent re-run)
    if 'EEU其他 TOTAL' in wb.sheetnames:
        wb.remove(wb['EEU其他 TOTAL'])

    # Pick formatting template: prefer EEU TOTAL, fall back to AZ
    template = 'EEU TOTAL' if 'EEU TOTAL' in wb.sheetnames else 'AZ'
    ws_new = wb.copy_worksheet(wb[template])
    ws_new.title = 'EEU其他 TOTAL'

    # Rename EEU → EEU（KZ、KG、BY）
    if 'EEU' in wb.sheetnames:
        wb['EEU'].title = 'EEU（KZ、KG、BY）'

    # Delete original EEU TOTAL
    if 'EEU TOTAL' in wb.sheetnames:
        wb.remove(wb['EEU TOTAL'])

    # Update header label and fill formulas
    ws_new['A1'] = 'CIS TOTAL'
    _fill_formulas(ws_new, _eeu_formula, wb['AZ'])

    # Move to position 0
    pos = wb.sheetnames.index('EEU其他 TOTAL')
    if pos > 0:
        wb.move_sheet('EEU其他 TOTAL', -pos)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def process_mena(file_bytes: bytes) -> bytes:
    """
    Idempotent MENA adjustment:
      - Remove existing 'SA TOTAL' if present
      - Copy 'MENA TOTAL' to get formatting template
      - Rename copy → 'SA TOTAL', update A1 → 'SA TOTAL'
      - Fill data rows with SA + SA_IOS formulas
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))

    # Remove stale SA TOTAL (idempotent re-run)
    if 'SA TOTAL' in wb.sheetnames:
        wb.remove(wb['SA TOTAL'])

    # Copy MENA TOTAL to get all formatting
    ws_new = wb.copy_worksheet(wb['MENA TOTAL'])
    ws_new.title = 'SA TOTAL'

    # Update header label and fill formulas
    ws_new['A1'] = 'SA TOTAL'
    _fill_formulas(ws_new, _sa_formula, wb['SA'])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
